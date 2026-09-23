import base64
import hashlib
import hmac
import io
import json
import os
import re
import secrets
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
import httpx
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from pypdf import PdfReader
from sqlalchemy import Boolean, DateTime, ForeignKey, Integer, String, Text, create_engine, inspect, or_, text
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, relationship, sessionmaker

from storage import create_storage

DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./lingyao.db")
engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
SECRET = os.getenv("SECRET_KEY", "dev-only-change-me").encode()
storage = create_storage(os.getenv("UPLOAD_DIR", "./uploads"))


class Base(DeclarativeBase): pass


class User(Base):
    __tablename__ = "users"
    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(80))
    phone: Mapped[str | None] = mapped_column(String(20), unique=True, index=True)
    email: Mapped[str] = mapped_column(String(160), unique=True, index=True)
    password_hash: Mapped[str] = mapped_column(String(200))
    role: Mapped[str] = mapped_column(String(30), default="consultant")

class RoleDefinition(Base):
    __tablename__ = "role_definitions"
    id: Mapped[int] = mapped_column(primary_key=True)
    code: Mapped[str] = mapped_column(String(40), unique=True, index=True)
    name: Mapped[str] = mapped_column(String(80), unique=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Candidate(Base):
    __tablename__ = "candidates"
    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(80), index=True)
    phone: Mapped[str | None] = mapped_column(String(40), index=True)
    email: Mapped[str | None] = mapped_column(String(160), index=True)
    city: Mapped[str | None] = mapped_column(String(80))
    current_company: Mapped[str | None] = mapped_column(String(160), index=True)
    current_title: Mapped[str | None] = mapped_column(String(160), index=True)
    school: Mapped[str | None] = mapped_column(String(160), index=True)
    degree: Mapped[str | None] = mapped_column(String(60))
    gender: Mapped[str | None] = mapped_column(String(20))
    graduation_time: Mapped[str | None] = mapped_column(String(40))
    years: Mapped[int | None] = mapped_column(Integer)
    expected_salary: Mapped[str | None] = mapped_column(String(80))
    summary: Mapped[str | None] = mapped_column(Text)
    visibility: Mapped[str] = mapped_column(String(20), default="company")
    owner_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Client(Base):
    __tablename__ = "clients"
    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(160), unique=True)
    industry: Mapped[str | None] = mapped_column(String(100))
    contact_name: Mapped[str | None] = mapped_column(String(80))
    contact_phone: Mapped[str | None] = mapped_column(String(40))
    owner_id: Mapped[int] = mapped_column(ForeignKey("users.id"))


class Job(Base):
    __tablename__ = "jobs"
    id: Mapped[int] = mapped_column(primary_key=True)
    title: Mapped[str] = mapped_column(String(160), index=True)
    client_id: Mapped[int] = mapped_column(ForeignKey("clients.id"))
    city: Mapped[str | None] = mapped_column(String(80))
    salary: Mapped[str | None] = mapped_column(String(80))
    status: Mapped[str] = mapped_column(String(30), default="open")
    owner_id: Mapped[int] = mapped_column(ForeignKey("users.id"))


class Application(Base):
    __tablename__ = "applications"
    id: Mapped[int] = mapped_column(primary_key=True)
    candidate_id: Mapped[int] = mapped_column(ForeignKey("candidates.id"))
    job_id: Mapped[int] = mapped_column(ForeignKey("jobs.id"))
    stage: Mapped[str] = mapped_column(String(30), default="搜寻")
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Note(Base):
    __tablename__ = "notes"
    id: Mapped[int] = mapped_column(primary_key=True)
    candidate_id: Mapped[int] = mapped_column(ForeignKey("candidates.id"))
    author_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    content: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Attachment(Base):
    __tablename__ = "attachments"
    id: Mapped[int] = mapped_column(primary_key=True)
    candidate_id: Mapped[int] = mapped_column(ForeignKey("candidates.id"))
    name: Mapped[str] = mapped_column(String(240))
    storage_key: Mapped[str] = mapped_column(String(240))
    content_type: Mapped[str | None] = mapped_column(String(120))
    size: Mapped[int] = mapped_column(Integer)
    parsed_text: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)

class AuditLog(Base):
    __tablename__ = "audit_logs"
    id: Mapped[int] = mapped_column(primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    action: Mapped[str] = mapped_column(String(80))
    entity_type: Mapped[str] = mapped_column(String(50))
    entity_id: Mapped[int | None] = mapped_column(Integer)
    detail: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)

class UserProfile(Base):
    __tablename__ = "user_profiles"
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"), primary_key=True)
    avatar_key: Mapped[str | None] = mapped_column(String(240))
    avatar_content_type: Mapped[str | None] = mapped_column(String(80))

class CompanySender(Base):
    __tablename__ = "company_senders"
    id: Mapped[int] = mapped_column(primary_key=True)
    client_id: Mapped[int | None] = mapped_column(ForeignKey("clients.id"), unique=True)
    display_name: Mapped[str] = mapped_column(String(160))
    form_header: Mapped[str] = mapped_column(String(240), default="求职申请表")
    interview_requirements: Mapped[str | None] = mapped_column(Text)
    sms_signature: Mapped[str | None] = mapped_column(String(80))
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)

class ResumeFlow(Base):
    __tablename__ = "resume_flows"
    id: Mapped[int] = mapped_column(primary_key=True)
    token: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    candidate_id: Mapped[int] = mapped_column(ForeignKey("candidates.id"))
    job_id: Mapped[int | None] = mapped_column(ForeignKey("jobs.id"))
    sender_id: Mapped[int | None] = mapped_column(ForeignKey("company_senders.id"))
    flow_type: Mapped[str] = mapped_column(String(30), default="application")
    recipient_phone: Mapped[str] = mapped_column(String(40))
    title: Mapped[str] = mapped_column(String(240))
    message: Mapped[str | None] = mapped_column(Text)
    form_data: Mapped[str | None] = mapped_column(Text)
    status: Mapped[str] = mapped_column(String(30), default="待填写")
    sms_status: Mapped[str] = mapped_column(String(30), default="链接已生成")
    created_by: Mapped[int] = mapped_column(ForeignKey("users.id"))
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    submitted_at: Mapped[datetime | None] = mapped_column(DateTime)
    expires_at: Mapped[datetime] = mapped_column(DateTime)

class FormTemplate(Base):
    __tablename__ = "form_templates"
    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(160))
    template_type: Mapped[str] = mapped_column(String(30), default="application")
    sender_id: Mapped[int | None] = mapped_column(ForeignKey("company_senders.id"))
    description: Mapped[str | None] = mapped_column(Text)
    fields_json: Mapped[str] = mapped_column(Text)
    source_file_name: Mapped[str | None] = mapped_column(String(240))
    created_by: Mapped[int] = mapped_column(ForeignKey("users.id"))
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class LoginIn(BaseModel): phone: str; password: str
class CandidateIn(BaseModel):
    name: str; phone: str | None = None; email: str | None = None; city: str | None = None
    current_company: str | None = None; current_title: str | None = None; school: str | None = None
    degree: str | None = None; gender: str | None = None; graduation_time: str | None = None; years: int | None = None; expected_salary: str | None = None
    summary: str | None = None; visibility: str = "company"; owner_id: int | None = None
class ClientIn(BaseModel): name: str; industry: str | None = None; contact_name: str | None = None; contact_phone: str | None = None
class JobIn(BaseModel): title: str; client_id: int; city: str | None = None; salary: str | None = None
class ApplicationIn(BaseModel): candidate_id: int; job_id: int; stage: str = "搜寻"
class StageIn(BaseModel): stage: str
class NoteIn(BaseModel): content: str
class UserIn(BaseModel): name: str; phone: str; email: str | None = None; password: str; role: str = "consultant"
class UserUpdate(BaseModel): name: str | None = None; phone: str | None = None; role: str | None = None; password: str | None = None
class RoleIn(BaseModel): name: str
class ClientUpdate(BaseModel): name: str; industry: str | None = None; contact_name: str | None = None; contact_phone: str | None = None
class JobUpdate(BaseModel): title: str; client_id: int; city: str | None = None; salary: str | None = None; status: str = "open"
class ProfileUpdate(BaseModel):
    name: str | None = None
    phone: str | None = None
    email: str | None = None
    old_password: str | None = None
    new_password: str | None = None
class SenderIn(BaseModel):
    client_id: int | None = None
    display_name: str
    form_header: str = "求职申请表"
    interview_requirements: str | None = None
    sms_signature: str | None = None
    enabled: bool = True
class ResumeFlowIn(BaseModel):
    candidate_id: int
    job_id: int | None = None
    sender_id: int | None = None
    flow_type: str = "application"
    recipient_phone: str
    title: str | None = None
    message: str | None = None
    template_id: int | None = None
class FormTemplateIn(BaseModel):
    name: str
    template_type: str = "application"
    sender_id: int | None = None
    description: str | None = None
    fields: list[dict]
class PublicFlowSubmit(BaseModel):
    answers: dict


def hash_password(password: str, salt: str = "lingyao") -> str:
    return hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 120_000).hex()


def valid_mobile(value: str) -> str:
    phone = clean_phone(value)
    if not phone or not re.fullmatch(r"1[3-9]\d{9}", phone): raise HTTPException(400, "请输入正确的11位手机号码")
    return phone


def valid_password(value: str) -> str:
    if not re.fullmatch(r"\d{6}", value): raise HTTPException(400, "密码必须为6位数字")
    return value


def token_for(user: User) -> str:
    payload = base64.urlsafe_b64encode(json.dumps({"id": user.id, "exp": int((datetime.now(timezone.utc)+timedelta(hours=12)).timestamp())}).encode()).decode().rstrip("=")
    return payload + "." + hmac.new(SECRET, payload.encode(), hashlib.sha256).hexdigest()


def get_db():
    with SessionLocal() as db: yield db


def current_user(authorization: str | None = Header(None), db: Session = Depends(get_db)) -> User:
    if not authorization or not authorization.startswith("Bearer "): raise HTTPException(401, "请先登录")
    token = authorization[7:]
    try:
        payload, sig = token.split(".")
        if not hmac.compare_digest(sig, hmac.new(SECRET, payload.encode(), hashlib.sha256).hexdigest()): raise ValueError()
        data = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload)%4)))
        if data["exp"] < datetime.now(timezone.utc).timestamp(): raise ValueError()
        user = db.get(User, data["id"])
        if not user or user.role in ("disabled", "deleted"): raise ValueError()
        return user
    except Exception: raise HTTPException(401, "登录已过期")


def clean_phone(v: str | None): return re.sub(r"\D", "", v or "") or None


def candidate_dict(c: Candidate, viewer: User, db: Session | None = None):
    can_contact = viewer.role in ("admin", "manager") or viewer.id == c.owner_id
    owner=db.get(User,c.owner_id) if db else None
    return {"id":c.id,"name":c.name,"phone":c.phone if can_contact else mask(c.phone),"email":c.email if can_contact else mask_email(c.email),"city":c.city,"current_company":c.current_company,"current_title":c.current_title,"school":c.school,"degree":c.degree,"gender":c.gender,"graduation_time":c.graduation_time,"years":c.years,"expected_salary":c.expected_salary,"summary":c.summary,"visibility":c.visibility,"owner_id":c.owner_id,"owner_name":owner.name if owner else f"顾问 #{c.owner_id}","can_view_contact":can_contact,"created_at":c.created_at}


def mask(v): return ("****" + v[-4:]) if v else None
def mask_email(v): return (v[:2] + "***@" + v.split("@",1)[1]) if v and "@" in v else None


def extract_resume(data: bytes, filename: str) -> tuple[str, dict]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf": text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    elif suffix == ".docx": text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    else: raise HTTPException(400, "仅支持 PDF 或 DOCX 简历")
    phone = re.search(r"(?<!\d)(1[3-9]\d{9})(?!\d)", text)
    email = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    return text, {"name": lines[0][:40] if lines else Path(filename).stem, "phone": phone.group(1) if phone else None, "email": email.group(0) if email else None, "summary": text[:3000]}


app = FastAPI(title="财富自由之路 API", version="1.0.0")
app.add_middleware(CORSMiddleware, allow_origins=os.getenv("CORS_ORIGINS", "http://localhost:3000,http://127.0.0.1:3000").split(","), allow_origin_regex=r"http://((10|192\.168)\.\d+\.\d+|172\.(1[6-9]|2\d|3[01])\.\d+\.\d+):3000", allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
def startup():
    Base.metadata.create_all(engine)
    if "phone" not in {c["name"] for c in inspect(engine).get_columns("users")}:
        with engine.begin() as conn: conn.execute(text("ALTER TABLE users ADD COLUMN phone VARCHAR(20)"))
    candidate_columns={c["name"] for c in inspect(engine).get_columns("candidates")}
    with engine.begin() as conn:
        if "gender" not in candidate_columns:conn.execute(text("ALTER TABLE candidates ADD COLUMN gender VARCHAR(20)"))
        if "graduation_time" not in candidate_columns:conn.execute(text("ALTER TABLE candidates ADD COLUMN graduation_time VARCHAR(40)"))
    with SessionLocal() as db:
        if not db.query(User).first(): seed(db)
        else:
            used={x.phone for x in db.query(User).filter(User.phone.is_not(None)).all()}
            for x in db.query(User).order_by(User.id).all():
                if not x.phone:
                    preferred=os.getenv("ADMIN_PHONE", "13800000000") if x.role=="admin" else f"137{x.id:08d}"
                    while preferred in used: preferred=str(int(preferred)+1)
                    x.phone=preferred;used.add(preferred);x.password_hash=hash_password("123456")
            db.commit()
        if not db.query(CompanySender).first():
            first_client=db.query(Client).first()
            db.add(CompanySender(client_id=first_client.id if first_client else None,display_name=first_client.name if first_client else "灵曜猎聘",form_header="求职申请表",interview_requirements="请提前10分钟到达，携带个人简历，并保持手机畅通。",sms_signature="灵曜猎聘"));db.commit()
        if not db.query(FormTemplate).first():
            owner=db.query(User).filter(User.role=="admin").first() or db.query(User).first()
            app_fields=[{"label":x,"type":"textarea" if x in ("离职原因及求职动机","主要工作经历") else "text","required":True} for x in ["姓名","手机号码","电子邮箱","现居住城市","最高学历","当前公司及职位","目前年薪","期望年薪","可到岗时间","离职原因及求职动机","主要工作经历"]]
            test_fields=[{"label":"我更喜欢的工作方式","type":"select","options":["独立完成","团队协作","两者均可"],"required":True},{"label":"面对压力时","type":"select","options":["立即行动解决","先分析再行动","寻求团队支持"],"required":True},{"label":"请描述您的优势","type":"textarea","required":True}]
            db.add_all([FormTemplate(name="标准求职申请表",template_type="application",fields_json=json.dumps(app_fields,ensure_ascii=False),created_by=owner.id),FormTemplate(name="基础性格测试",template_type="personality",fields_json=json.dumps(test_fields,ensure_ascii=False),created_by=owner.id)]);db.commit()


@app.get("/health")
def health(): return {"status":"ok"}


@app.post("/api/auth/login")
def login(body: LoginIn, db: Session = Depends(get_db)):
    phone=valid_mobile(body.phone);valid_password(body.password)
    user = db.query(User).filter(User.phone == phone).first()
    if not user or user.role in ("disabled", "deleted") or user.password_hash != hash_password(body.password): raise HTTPException(401, "账号或密码错误")
    return {"access_token":token_for(user),"user":{"id":user.id,"name":user.name,"phone":user.phone,"email":user.email,"role":user.role}}


@app.get("/api/me")
def me(db:Session=Depends(get_db),user: User = Depends(current_user)):
    profile=db.get(UserProfile,user.id)
    return {"id":user.id,"name":user.name,"phone":user.phone,"email":user.email,"role":user.role,"avatar_url":f"/api/avatars/{profile.avatar_key}" if profile and profile.avatar_key else None}

@app.patch("/api/me")
def update_me(body:ProfileUpdate,db:Session=Depends(get_db),user:User=Depends(current_user)):
    if body.name:
        if len(body.name.strip())<2:raise HTTPException(400,"姓名至少 2 个字符")
        user.name=body.name.strip()
    if body.phone:
        phone=valid_mobile(body.phone)
        if db.query(User).filter(User.phone==phone,User.id!=user.id).first():raise HTTPException(409,"该手机号已被使用")
        user.phone=phone
    if body.email:
        email=body.email.lower().strip()
        duplicate=db.query(User).filter(User.email==email,User.id!=user.id).first()
        if duplicate:raise HTTPException(409,"该邮箱已被使用")
        user.email=email
    if body.new_password:
        if not body.old_password or user.password_hash!=hash_password(body.old_password):raise HTTPException(400,"当前密码不正确")
        user.password_hash=hash_password(valid_password(body.new_password))
    audit(db,user,"update","profile",user.id,user.email);db.commit()
    return me(db,user)

@app.post("/api/me/avatar")
async def upload_avatar(file:UploadFile=File(...),db:Session=Depends(get_db),user:User=Depends(current_user)):
    if file.content_type not in ("image/jpeg","image/png","image/webp"):raise HTTPException(400,"头像仅支持 JPG、PNG 或 WebP")
    data=await file.read()
    if len(data)>3*1024*1024:raise HTTPException(413,"头像不能超过 3MB")
    key,_=storage.save(data,file.filename or "avatar.jpg")
    profile=db.get(UserProfile,user.id) or UserProfile(user_id=user.id)
    profile.avatar_key=key;profile.avatar_content_type=file.content_type;db.add(profile);audit(db,user,"upload","avatar",user.id,key);db.commit()
    return {"avatar_url":f"/api/avatars/{key}"}

@app.get("/api/avatars/{key}")
def avatar(key:str,db:Session=Depends(get_db)):
    profile=db.query(UserProfile).filter(UserProfile.avatar_key==Path(key).name).first()
    if not profile:raise HTTPException(404,"头像不存在")
    return FileResponse(storage.path(key),media_type=profile.avatar_content_type)


@app.get("/api/candidates")
def candidates(q: str="", city: str|None=None, company: str|None=None, school: str|None=None, years_min: int|None=None, scope: str="company", db: Session=Depends(get_db), user: User=Depends(current_user)):
    query=db.query(Candidate)
    if scope=="mine": query=query.filter(Candidate.owner_id==user.id)
    if q:
        like=f"%{q}%"; query=query.filter(or_(Candidate.name.ilike(like),Candidate.current_company.ilike(like),Candidate.current_title.ilike(like),Candidate.school.ilike(like),Candidate.summary.ilike(like)))
    if city: query=query.filter(Candidate.city==city)
    if company: query=query.filter(Candidate.current_company.ilike(f"%{company}%"))
    if school: query=query.filter(Candidate.school.ilike(f"%{school}%"))
    if years_min is not None: query=query.filter(Candidate.years>=years_min)
    return [candidate_dict(c,user,db) for c in query.order_by(Candidate.created_at.desc()).all()]


@app.post("/api/candidates", status_code=201)
def create_candidate(body: CandidateIn, db: Session=Depends(get_db), user: User=Depends(current_user)):
    duplicate = find_duplicate(db, body.phone, body.email)
    if duplicate: raise HTTPException(409, {"message":"发现重复人才","candidate_id":duplicate.id,"name":duplicate.name})
    data=body.model_dump(); data["phone"]=clean_phone(data["phone"]); data["owner_id"]=data["owner_id"] or user.id
    c=Candidate(**data); db.add(c); db.commit(); return candidate_dict(c,user,db)


@app.put("/api/candidates/{candidate_id}")
def update_candidate(candidate_id:int, body:CandidateIn, db:Session=Depends(get_db), user:User=Depends(current_user)):
    c=db.get(Candidate,candidate_id)
    if not c: raise HTTPException(404,"候选人不存在")
    if user.role not in ("admin","manager") and c.owner_id!=user.id: raise HTTPException(403,"无权编辑该人才")
    for k,v in body.model_dump(exclude_unset=True).items(): setattr(c,k,clean_phone(v) if k=="phone" else v)
    db.commit(); return candidate_dict(c,user,db)


@app.get("/api/candidates/duplicates/check")
def duplicates(phone:str|None=None,email:str|None=None,db:Session=Depends(get_db),user:User=Depends(current_user)):
    c=find_duplicate(db,phone,email); return {"duplicate":bool(c),"candidate":candidate_dict(c,user,db) if c else None}


def find_duplicate(db,phone,email):
    conditions=[]
    if clean_phone(phone): conditions.append(Candidate.phone==clean_phone(phone))
    if email: conditions.append(Candidate.email==email.lower().strip())
    return db.query(Candidate).filter(or_(*conditions)).first() if conditions else None

def audit(db: Session, user: User, action: str, entity_type: str, entity_id: int | None = None, detail: str | None = None):
    db.add(AuditLog(user_id=user.id, action=action, entity_type=entity_type, entity_id=entity_id, detail=detail))

def require_admin(user: User):
    if user.role not in ("admin", "manager"): raise HTTPException(403, "仅管理员可执行此操作")

BUILTIN_ROLES = {"admin":"管理员", "manager":"经理", "consultant":"顾问"}

def valid_role(db: Session, role: str) -> bool:
    return role in BUILTIN_ROLES or db.query(RoleDefinition).filter(RoleDefinition.code == role).first() is not None

def flow_dict(x: ResumeFlow, db: Session, include_answers: bool = True):
    candidate=db.get(Candidate,x.candidate_id);job=db.get(Job,x.job_id) if x.job_id else None;sender=db.get(CompanySender,x.sender_id) if x.sender_id else None
    payload=json.loads(x.message) if x.message and x.message.startswith('{"template_id"') else {}
    template=db.get(FormTemplate,payload.get("template_id")) if payload.get("template_id") else None
    return {"id":x.id,"token":x.token,"candidate_id":x.candidate_id,"candidate_name":candidate.name if candidate else "—","job_id":x.job_id,"job_title":job.title if job else None,"sender_id":x.sender_id,"sender_name":sender.display_name if sender else "灵曜猎聘","form_header":sender.form_header if sender else "求职申请表","interview_requirements":sender.interview_requirements if sender else None,"flow_type":x.flow_type,"recipient_phone":x.recipient_phone,"title":x.title,"message":payload.get("message",x.message),"template_id":template.id if template else None,"template_name":template.name if template else None,"fields":json.loads(template.fields_json) if template else None,"status":x.status,"sms_status":x.sms_status,"answers":json.loads(x.form_data) if include_answers and x.form_data else None,"created_at":x.created_at,"submitted_at":x.submitted_at,"expires_at":x.expires_at,"public_url":f'{os.getenv("PUBLIC_WEB_URL","http://localhost:3000")}/?form={x.token}'}


@app.post("/api/resumes/parse")
async def parse_resume(file:UploadFile=File(...), user:User=Depends(current_user)):
    data=await file.read()
    if len(data)>15*1024*1024: raise HTTPException(413,"文件不能超过 15MB")
    text, fields=extract_resume(data,file.filename or "resume.pdf")
    return {"fields":fields,"text_preview":text[:800]}


@app.post("/api/resumes/import", status_code=201)
async def import_resume(file:UploadFile=File(...), job_id:int|None=Form(None), db:Session=Depends(get_db), user:User=Depends(current_user)):
    data=await file.read()
    if len(data)>15*1024*1024: raise HTTPException(413,"文件不能超过 15MB")
    filename=file.filename or "resume.pdf"
    text,fields=extract_resume(data,filename)
    candidate=find_duplicate(db,fields.get("phone"),fields.get("email"))
    created=False
    if not candidate:
        candidate=Candidate(name=fields.get("name") or Path(filename).stem,phone=clean_phone(fields.get("phone")),email=(fields.get("email") or None),summary=fields.get("summary"),owner_id=user.id,visibility="company")
        db.add(candidate);db.flush();created=True
    key,size=storage.save(data,filename)
    db.add(Attachment(candidate_id=candidate.id,name=filename,storage_key=key,content_type=file.content_type,size=size,parsed_text=text))
    if job_id:
        if not db.get(Job,job_id):raise HTTPException(404,"岗位分组不存在")
        if not db.query(Application).filter(Application.candidate_id==candidate.id,Application.job_id==job_id).first():
            db.add(Application(candidate_id=candidate.id,job_id=job_id,stage="搜寻"))
    audit(db,user,"import","resume",candidate.id,filename);db.commit()
    return {"candidate":candidate_dict(candidate,user,db),"created":created,"filename":filename}


@app.post("/api/candidates/{candidate_id}/attachments", status_code=201)
async def attach(candidate_id:int,file:UploadFile=File(...),db:Session=Depends(get_db),user:User=Depends(current_user)):
    if not db.get(Candidate,candidate_id): raise HTTPException(404,"候选人不存在")
    data=await file.read(); text=None
    if len(data)>15*1024*1024: raise HTTPException(413,"文件不能超过 15MB")
    if Path(file.filename or "").suffix.lower() in (".pdf",".docx"):
        text,_=extract_resume(data,file.filename or "resume.pdf")
    key,size=storage.save(data,file.filename or "attachment")
    a=Attachment(candidate_id=candidate_id,name=file.filename or "attachment",storage_key=key,content_type=file.content_type,size=size,parsed_text=text);db.add(a);db.commit()
    return {"id":a.id,"name":a.name,"size":a.size}


@app.get("/api/attachments/{attachment_id}")
def download(attachment_id:int,db:Session=Depends(get_db),user:User=Depends(current_user)):
    a=db.get(Attachment,attachment_id)
    if not a: raise HTTPException(404,"附件不存在")
    return FileResponse(storage.path(a.storage_key),filename=a.name,media_type=a.content_type)


@app.get("/api/clients")
def clients(db:Session=Depends(get_db),user:User=Depends(current_user)): return db.query(Client).order_by(Client.name).all()
@app.post("/api/clients",status_code=201)
def create_client(body:ClientIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    c=Client(**body.model_dump(),owner_id=user.id);db.add(c);db.commit();return c
@app.get("/api/jobs")
def jobs(status:str|None=None,db:Session=Depends(get_db),user:User=Depends(current_user)):
    q=db.query(Job);return (q.filter(Job.status==status) if status else q).order_by(Job.id.desc()).all()
@app.post("/api/jobs",status_code=201)
def create_job(body:JobIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    if not db.get(Client,body.client_id): raise HTTPException(404,"客户不存在")
    j=Job(**body.model_dump(),owner_id=user.id);db.add(j);db.commit();return j

@app.get("/api/resume-flow/senders")
def list_senders(db:Session=Depends(get_db),user:User=Depends(current_user)):
    return db.query(CompanySender).order_by(CompanySender.id).all()

@app.post("/api/resume-flow/senders",status_code=201)
def create_sender(body:SenderIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user)
    if body.client_id and not db.get(Client,body.client_id):raise HTTPException(404,"客户不存在")
    if body.client_id and db.query(CompanySender).filter(CompanySender.client_id==body.client_id).first():raise HTTPException(409,"该客户已配置发送主体")
    x=CompanySender(**body.model_dump());db.add(x);db.flush();audit(db,user,"create","company_sender",x.id,x.display_name);db.commit();return x

@app.put("/api/resume-flow/senders/{sender_id}")
def update_sender(sender_id:int,body:SenderIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user);x=db.get(CompanySender,sender_id)
    if not x:raise HTTPException(404,"发送主体不存在")
    for k,v in body.model_dump().items():setattr(x,k,v)
    audit(db,user,"update","company_sender",x.id,x.display_name);db.commit();return x

@app.get("/api/resume-flow/templates")
def list_form_templates(db:Session=Depends(get_db),user:User=Depends(current_user)):
    return [{"id":x.id,"name":x.name,"template_type":x.template_type,"sender_id":x.sender_id,"description":x.description,"fields":json.loads(x.fields_json),"source_file_name":x.source_file_name,"created_at":x.created_at} for x in db.query(FormTemplate).order_by(FormTemplate.id.desc()).all()]

@app.post("/api/resume-flow/templates",status_code=201)
def create_form_template(body:FormTemplateIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    if body.template_type not in ("application","personality"):raise HTTPException(400,"模板类型无效")
    fields=[f for f in body.fields if str(f.get("label","")).strip()]
    if not fields:raise HTTPException(400,"请至少添加一个表单问题")
    x=FormTemplate(name=body.name,template_type=body.template_type,sender_id=body.sender_id,description=body.description,fields_json=json.dumps(fields,ensure_ascii=False),created_by=user.id);db.add(x);db.commit();return {"id":x.id,"name":x.name,"fields":fields}

@app.put("/api/resume-flow/templates/{template_id}")
def update_form_template(template_id:int,body:FormTemplateIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    x=db.get(FormTemplate,template_id)
    if not x:raise HTTPException(404,"模板不存在")
    if user.role not in ("admin","manager") and x.created_by!=user.id:raise HTTPException(403,"无权修改该模板")
    if body.template_type not in ("application","personality"):raise HTTPException(400,"模板类型无效")
    fields=[f for f in body.fields if str(f.get("label","")).strip()]
    if not fields:raise HTTPException(400,"请至少保留一个表单问题")
    x.name=body.name;x.template_type=body.template_type;x.sender_id=body.sender_id;x.description=body.description;x.fields_json=json.dumps(fields,ensure_ascii=False)
    audit(db,user,"update","form_template",x.id,x.name);db.commit()
    return {"id":x.id,"name":x.name,"fields":fields}

@app.delete("/api/resume-flow/templates/{template_id}",status_code=204)
def delete_form_template(template_id:int,db:Session=Depends(get_db),user:User=Depends(current_user)):
    x=db.get(FormTemplate,template_id)
    if not x:raise HTTPException(404,"模板不存在")
    if user.role not in ("admin","manager") and x.created_by!=user.id:raise HTTPException(403,"无权删除该模板")
    audit(db,user,"delete","form_template",x.id,x.name);db.delete(x);db.commit()

@app.post("/api/resume-flow/templates/upload",status_code=201)
async def upload_form_template(name:str=Form(...),template_type:str=Form("application"),sender_id:int|None=Form(None),file:UploadFile=File(...),db:Session=Depends(get_db),user:User=Depends(current_user)):
    data=await file.read()
    if len(data)>10*1024*1024:raise HTTPException(413,"模板文件不能超过10MB")
    suffix=Path(file.filename or "").suffix.lower()
    if suffix==".docx":lines=[p.text.strip() for p in Document(io.BytesIO(data)).paragraphs if p.text.strip()]
    elif suffix==".pdf":lines=[line.strip() for p in PdfReader(io.BytesIO(data)).pages for line in (p.extract_text() or "").splitlines() if line.strip()]
    elif suffix==".txt":lines=[line.strip() for line in data.decode("utf-8-sig").splitlines() if line.strip()]
    else:raise HTTPException(400,"仅支持 DOCX、PDF 或 TXT 模板")
    fields=[{"label":line[:120],"type":"textarea" if len(line)>35 else "text","required":True} for line in lines[:60]]
    if not fields:raise HTTPException(400,"文件中未识别到可用文字")
    x=FormTemplate(name=name,template_type=template_type,sender_id=sender_id,fields_json=json.dumps(fields,ensure_ascii=False),source_file_name=file.filename,created_by=user.id);db.add(x);db.commit();return {"id":x.id,"name":x.name,"fields":fields,"source_file_name":x.source_file_name}

@app.get("/api/resume-flow")
def list_resume_flows(db:Session=Depends(get_db),user:User=Depends(current_user)):
    return [flow_dict(x,db) for x in db.query(ResumeFlow).order_by(ResumeFlow.id.desc()).all()]

@app.post("/api/resume-flow",status_code=201)
def create_resume_flow(body:ResumeFlowIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    if body.flow_type not in ("application","personality","interview"):raise HTTPException(400,"无效流程类型")
    candidate=db.get(Candidate,body.candidate_id)
    if not candidate:raise HTTPException(404,"候选人不存在")
    if body.job_id and not db.get(Job,body.job_id):raise HTTPException(404,"职位不存在")
    sender=db.get(CompanySender,body.sender_id) if body.sender_id else None
    if body.sender_id and not sender:raise HTTPException(404,"发送主体不存在")
    phone=clean_phone(body.recipient_phone)
    if not phone:raise HTTPException(400,"请填写候选人手机号")
    labels={"application":"求职申请表","personality":"性格测试","interview":"面试邀请"}
    template=db.get(FormTemplate,body.template_id) if body.template_id else None
    if body.flow_type!="interview" and not template:raise HTTPException(400,"请选择申请表或性格测试模板")
    if template and template.template_type!=body.flow_type:raise HTTPException(400,"模板类型与发送内容不匹配")
    stored_message=json.dumps({"template_id":template.id,"message":body.message},ensure_ascii=False) if template else body.message
    x=ResumeFlow(token=secrets.token_urlsafe(24),candidate_id=candidate.id,job_id=body.job_id,sender_id=body.sender_id,flow_type=body.flow_type,recipient_phone=phone,title=body.title or f"{sender.display_name if sender else '灵曜猎聘'} · {labels[body.flow_type]}",message=stored_message,status="已发送" if body.flow_type=="interview" else "待填写",created_by=user.id,expires_at=datetime.utcnow()+timedelta(days=14))
    db.add(x);db.flush()
    url=f'{os.getenv("PUBLIC_WEB_URL","http://localhost:3000")}/?form={x.token}'
    webhook=os.getenv("SMS_WEBHOOK_URL")
    if webhook:
        try:
            r=httpx.post(webhook,json={"phone":phone,"signature":sender.sms_signature if sender else "灵曜猎聘","content":f"{x.title}：{url}"},timeout=10)
            r.raise_for_status();x.sms_status="短信已发送"
        except Exception:x.sms_status="短信发送失败"
    else:x.sms_status="演示模式·请复制链接"
    audit(db,user,"send",f"resume_flow_{body.flow_type}",x.id,phone);db.commit();return flow_dict(x,db)

@app.get("/api/public/resume-flow/{token}")
def public_resume_flow(token:str,db:Session=Depends(get_db)):
    x=db.query(ResumeFlow).filter(ResumeFlow.token==token).first()
    if not x:raise HTTPException(404,"链接不存在")
    if x.expires_at<datetime.utcnow():raise HTTPException(410,"链接已过期")
    result=flow_dict(x,db,False);result.pop("recipient_phone",None);result.pop("public_url",None);return result

@app.post("/api/public/resume-flow/{token}")
def submit_public_resume_flow(token:str,body:PublicFlowSubmit,db:Session=Depends(get_db)):
    x=db.query(ResumeFlow).filter(ResumeFlow.token==token).first()
    if not x:raise HTTPException(404,"链接不存在")
    if x.expires_at<datetime.utcnow():raise HTTPException(410,"链接已过期")
    if x.flow_type=="interview":raise HTTPException(400,"面试邀请无需填写")
    x.form_data=json.dumps(body.answers,ensure_ascii=False);x.status="已填写";x.submitted_at=datetime.utcnow();db.commit();return {"ok":True,"message":"提交成功"}


STAGES=["搜寻","联系","推荐","面试","Offer","入职","保证期","回款"]
@app.get("/api/pipeline")
def pipeline(job_id:int|None=None,db:Session=Depends(get_db),user:User=Depends(current_user)):
    q=db.query(Application); rows=(q.filter(Application.job_id==job_id) if job_id else q).all()
    return [{"id":x.id,"candidate_id":x.candidate_id,"job_id":x.job_id,"stage":x.stage,"updated_at":x.updated_at} for x in rows]
@app.post("/api/pipeline",status_code=201)
def create_application(body:ApplicationIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    if body.stage not in STAGES: raise HTTPException(400,"无效阶段")
    if not db.get(Candidate,body.candidate_id) or not db.get(Job,body.job_id): raise HTTPException(404,"人才或职位不存在")
    a=Application(**body.model_dump());db.add(a);db.commit();return a
@app.patch("/api/pipeline/{application_id}/stage")
def move_stage(application_id:int,body:StageIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    if body.stage not in STAGES: raise HTTPException(400,"无效阶段")
    a=db.get(Application,application_id)
    if not a: raise HTTPException(404,"流程不存在")
    a.stage=body.stage;a.updated_at=datetime.utcnow();db.commit();return a
@app.get("/api/candidates/{candidate_id}/notes")
def notes(candidate_id:int,db:Session=Depends(get_db),user:User=Depends(current_user)): return db.query(Note).filter(Note.candidate_id==candidate_id).order_by(Note.id.desc()).all()
@app.post("/api/candidates/{candidate_id}/notes",status_code=201)
def create_note(candidate_id:int,body:NoteIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    if not db.get(Candidate,candidate_id): raise HTTPException(404,"候选人不存在")
    n=Note(candidate_id=candidate_id,author_id=user.id,content=body.content);db.add(n);db.commit();return n

@app.get("/api/candidates/{candidate_id}")
def candidate_detail(candidate_id:int, db:Session=Depends(get_db), user:User=Depends(current_user)):
    c=db.get(Candidate,candidate_id)
    if not c: raise HTTPException(404,"候选人不存在")
    result=candidate_dict(c,user,db)
    result["notes"]=[{"id":n.id,"content":n.content,"author_id":n.author_id,"created_at":n.created_at} for n in db.query(Note).filter(Note.candidate_id==candidate_id).order_by(Note.id.desc()).all()]
    result["attachments"]=[{"id":a.id,"name":a.name,"size":a.size,"content_type":a.content_type,"created_at":a.created_at} for a in db.query(Attachment).filter(Attachment.candidate_id==candidate_id).order_by(Attachment.id.desc()).all()]
    result["applications"]=[{"id":a.id,"job_id":a.job_id,"stage":a.stage,"updated_at":a.updated_at} for a in db.query(Application).filter(Application.candidate_id==candidate_id).all()]
    return result

@app.delete("/api/candidates/{candidate_id}",status_code=204)
def delete_candidate(candidate_id:int,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user);c=db.get(Candidate,candidate_id)
    if not c:raise HTTPException(404,"候选人不存在")
    for model in (Note,Attachment,Application):db.query(model).filter(model.candidate_id==candidate_id).delete()
    audit(db,user,"delete","candidate",candidate_id,c.name);db.delete(c);db.commit()

@app.put("/api/clients/{client_id}")
def update_client(client_id:int,body:ClientUpdate,db:Session=Depends(get_db),user:User=Depends(current_user)):
    c=db.get(Client,client_id)
    if not c:raise HTTPException(404,"客户不存在")
    for k,v in body.model_dump().items():setattr(c,k,v)
    audit(db,user,"update","client",c.id,c.name);db.commit();return c

@app.put("/api/jobs/{job_id}")
def update_job(job_id:int,body:JobUpdate,db:Session=Depends(get_db),user:User=Depends(current_user)):
    j=db.get(Job,job_id)
    if not j:raise HTTPException(404,"职位不存在")
    for k,v in body.model_dump().items():setattr(j,k,v)
    audit(db,user,"update","job",j.id,j.title);db.commit();return j

@app.get("/api/dashboard")
def dashboard(db:Session=Depends(get_db),user:User=Depends(current_user)):
    return {"candidates":db.query(Candidate).count(),"clients":db.query(Client).count(),"open_jobs":db.query(Job).filter(Job.status=="open").count(),"active_pipeline":db.query(Application).filter(Application.stage!="回款").count(),"stages":{stage:db.query(Application).filter(Application.stage==stage).count() for stage in STAGES}}

@app.get("/api/users")
def users(db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user);return [{"id":x.id,"name":x.name,"phone":x.phone,"email":x.email,"role":x.role} for x in db.query(User).filter(User.role!="deleted").order_by(User.id).all()]

@app.get("/api/roles")
def list_roles(db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user)
    builtins=[{"code":code,"name":name,"builtin":True} for code,name in BUILTIN_ROLES.items()]
    custom=[{"code":x.code,"name":x.name,"builtin":False} for x in db.query(RoleDefinition).order_by(RoleDefinition.id).all()]
    return builtins+custom

@app.post("/api/roles",status_code=201)
def create_role(body:RoleIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user);name=body.name.strip()
    if not name:raise HTTPException(400,"请输入角色名称")
    if name in BUILTIN_ROLES.values() or db.query(RoleDefinition).filter(RoleDefinition.name==name).first():raise HTTPException(409,"角色名称已存在")
    x=RoleDefinition(code=f"custom_{secrets.token_hex(4)}",name=name);db.add(x);db.flush();audit(db,user,"create","role",x.id,name);db.commit()
    return {"code":x.code,"name":x.name,"builtin":False}

@app.delete("/api/roles/{code}",status_code=204)
def delete_role(code:str,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user)
    if code in BUILTIN_ROLES:raise HTTPException(400,"系统内置角色不能删除")
    x=db.query(RoleDefinition).filter(RoleDefinition.code==code).first()
    if not x:raise HTTPException(404,"角色不存在")
    if db.query(User).filter(User.role==code).first():raise HTTPException(409,"请先将使用该角色的成员调整为其他角色")
    audit(db,user,"delete","role",x.id,x.name);db.delete(x);db.commit()

@app.post("/api/users",status_code=201)
def create_user(body:UserIn,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user)
    if not valid_role(db,body.role):raise HTTPException(400,"无效角色")
    phone=valid_mobile(body.phone);valid_password(body.password)
    if db.query(User).filter(User.phone==phone).first():raise HTTPException(409,"手机号已存在")
    email=(body.email or f"{phone}@lingyao.local").lower()
    if db.query(User).filter(User.email==email).first():raise HTTPException(409,"邮箱已存在")
    x=User(name=body.name,phone=phone,email=email,password_hash=hash_password(body.password),role=body.role);db.add(x);db.flush();audit(db,user,"create","user",x.id,x.phone);db.commit();return {"id":x.id,"name":x.name,"phone":x.phone,"email":x.email,"role":x.role}

@app.patch("/api/users/{user_id}")
def update_user(user_id:int,body:UserUpdate,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user);x=db.get(User,user_id)
    if not x:raise HTTPException(404,"用户不存在")
    if body.name:x.name=body.name
    if body.phone:
        phone=valid_mobile(body.phone)
        if db.query(User).filter(User.phone==phone,User.id!=x.id).first():raise HTTPException(409,"手机号已存在")
        x.phone=phone
    if body.role:
        if body.role!="disabled" and not valid_role(db,body.role):raise HTTPException(400,"无效角色")
        if x.id==user.id and body.role=="disabled":raise HTTPException(400,"不能停用当前账号")
        x.role=body.role
    if body.password:
        x.password_hash=hash_password(valid_password(body.password))
    audit(db,user,"update","user",x.id,x.phone);db.commit();return {"id":x.id,"name":x.name,"phone":x.phone,"email":x.email,"role":x.role}

@app.delete("/api/users/{user_id}",status_code=204)
def delete_user(user_id:int,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user);x=db.get(User,user_id)
    if not x or x.role=="deleted":raise HTTPException(404,"成员不存在")
    if x.id==user.id:raise HTTPException(400,"不能删除当前登录账号")
    old_name=x.name;audit(db,user,"delete","user",x.id,old_name)
    x.name=f"已删除成员 #{x.id}";x.phone=None;x.email=f"deleted-{x.id}-{secrets.token_hex(4)}@invalid.local";x.password_hash=hash_password(secrets.token_hex(16));x.role="deleted";db.commit()

@app.get("/api/audit-logs")
def audit_logs(limit:int=100,db:Session=Depends(get_db),user:User=Depends(current_user)):
    require_admin(user);rows=db.query(AuditLog).order_by(AuditLog.id.desc()).limit(min(limit,500)).all()
    return [{"id":x.id,"user_id":x.user_id,"action":x.action,"entity_type":x.entity_type,"entity_id":x.entity_id,"detail":x.detail,"created_at":x.created_at} for x in rows]


def seed(db:Session):
    admin=User(name="林若安",phone=os.getenv("ADMIN_PHONE", "13800000000"),email="admin@lingyao.local",password_hash=hash_password("123456"),role="admin")
    consultant=User(name="王晨",phone="13900000000",email="consultant@lingyao.local",password_hash=hash_password("123456"),role="consultant")
    db.add_all([admin,consultant]);db.flush()
    people=[Candidate(name="陈思远",phone="13800138001",email="siyuan@example.com",city="上海",current_company="云岚科技",current_title="高级算法工程师",school="上海交通大学",degree="硕士",years=8,owner_id=admin.id,summary="推荐系统、NLP 与算法团队管理经验"),Candidate(name="周雨桐",phone="13800138002",email="yutong@example.com",city="苏州",current_company="星河智能",current_title="海外市场总监",school="南京大学",degree="本科",years=10,owner_id=consultant.id,summary="制造业出海与欧美市场拓展"),Candidate(name="许文博",phone="13800138003",email="wenbo@example.com",city="杭州",current_company="澄明医疗",current_title="供应链负责人",school="浙江大学",degree="MBA",years=12,owner_id=admin.id,summary="医疗器械供应链体系搭建")]
    client=Client(name="云岚科技",industry="企业服务 / AI",contact_name="张总",contact_phone="13900001111",owner_id=admin.id)
    db.add_all(people+[client]);db.flush()
    job=Job(title="算法平台负责人",client_id=client.id,city="上海",salary="60-90万",owner_id=admin.id);db.add(job);db.flush()
    db.add_all([Application(candidate_id=people[0].id,job_id=job.id,stage="推荐"),Application(candidate_id=people[1].id,job_id=job.id,stage="面试"),Note(candidate_id=people[0].id,author_id=admin.id,content="候选人意向积极，本周可安排客户面试。")]);db.commit()
